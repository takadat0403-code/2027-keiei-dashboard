from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "source" / "2027年度_経営改革ダッシュボード_土台.xlsx"
OUTPUT = ROOT / "src" / "data" / "dashboard.json"


def iso(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def numeric_or_none(value: Any) -> float | int | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    return None


def read_table(ws, start_row: int, columns: list[str], stop_when_blank_id: bool = False) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row_idx in range(start_row, ws.max_row + 1):
        values = [ws.cell(row_idx, col + 1).value for col in range(len(columns))]
        if stop_when_blank_id and values[0] in (None, ""):
            continue
        if not any(v not in (None, "") and not (isinstance(v, str) and v.startswith("=")) for v in values):
            continue
        row: dict[str, Any] = {}
        for key, value in zip(columns, values):
            if isinstance(value, str) and value.startswith("="):
                row[key] = None
            else:
                row[key] = iso(value)
        rows.append(row)
    return rows


def main() -> None:
    wb = load_workbook(SOURCE, data_only=False)

    master_ws = wb["Master"]
    master = {}
    for r in range(5, 15):
        key = master_ws.cell(r, 1).value
        if key:
            master[key] = iso(master_ws.cell(r, 2).value)

    prices = []
    for r in range(19, 23):
        prices.append({
            "facility": master_ws.cell(r, 1).value,
            "course": master_ws.cell(r, 2).value,
            "style": master_ws.cell(r, 3).value,
            "weekdayPrice": numeric_or_none(master_ws.cell(r, 4).value),
            "weekendPrice": numeric_or_none(master_ws.cell(r, 5).value),
            "positioning": master_ws.cell(r, 6).value,
            "note": master_ws.cell(r, 7).value,
        })

    revenue_columns = [
        "id", "year", "month", "facility", "course", "style", "dayType", "timeBand",
        "basePrice", "plannedPrice", "actualAvgPrice", "capacitySlots", "bookedSlots", "bookingRate",
        "plannedPlayers", "actualPlayers", "plannedRevenue", "actualRevenue", "priceStage", "note"
    ]
    revenue = read_table(wb["料金収益"], 5, revenue_columns, stop_when_blank_id=True)
    thresholds = {
        "discountBelow": float(master["予約率_割引検討上限"]),
        "standardBelow": float(master["予約率_基本販売上限"]),
        "premiumFrom": float(master["予約率_プレミアム検討"]),
    }
    for row in revenue:
        if row["capacitySlots"] is not None and row["bookedSlots"] is not None and row["capacitySlots"]:
            row["bookingRate"] = row["bookedSlots"] / row["capacitySlots"]
        else:
            row["bookingRate"] = None
        row["plannedRevenue"] = (
            row["plannedPrice"] * row["plannedPlayers"]
            if row["plannedPrice"] is not None and row["plannedPlayers"] is not None else None
        )
        row["actualRevenue"] = (
            row["actualAvgPrice"] * row["actualPlayers"]
            if row["actualAvgPrice"] is not None and row["actualPlayers"] is not None else None
        )
        rate = row["bookingRate"]
        if rate is None:
            row["priceStage"] = "未入力"
        elif rate < thresholds["discountBelow"]:
            row["priceStage"] = "割引検討"
        elif rate < thresholds["standardBelow"]:
            row["priceStage"] = "基本販売"
        elif rate < thresholds["premiumFrom"]:
            row["priceStage"] = "割引停止"
        else:
            row["priceStage"] = "プレミアム価格検討"

    course_columns = [
        "date", "facility", "course", "green", "turfDensity", "rootLengthMm", "thatchMm", "rootZoneMoisturePct",
        "permeabilityMmH", "soilHardness", "soilTempC", "surfaceTempC", "irrigationUniformityPct", "diseaseScore",
        "wateringHours", "chemFertilizerCost", "renovationHours", "weaknessFlag", "owner", "note", "measurementSource"
    ]
    course_quality = read_table(wb["コース品質"], 5, course_columns)

    workforce_columns = [
        "date", "facility", "department", "plannedHeadcount", "actualHeadcount", "headcountVariance", "laborHours",
        "overtimeHours", "paidLeavePeople", "clockCorrections", "throughput", "productivity", "laborCost", "owner", "note", "source"
    ]
    workforce = read_table(wb["労務生産性"], 5, workforce_columns)

    dx_columns = [
        "id", "category", "initiative", "purpose", "owner", "startDate", "dueDate", "priority", "status", "progress",
        "kpi", "actualState", "nextAction", "risk", "note"
    ]
    dx = read_table(wb["DX_CRM"], 5, dx_columns, stop_when_blank_id=True)

    investment_columns = [
        "id", "facility", "category", "project", "purpose", "estimatedCost", "subsidy", "netCost", "savedHours", "hourlyCost",
        "laborSavings", "otherAnnualBenefit", "annualBenefit", "paybackYears", "qualityImpact", "bcpImpact", "priority", "note"
    ]
    investments = read_table(wb["投資ROI"], 5, investment_columns, stop_when_blank_id=True)
    for row in investments:
        row["netCost"] = (
            row["estimatedCost"] - (row["subsidy"] or 0)
            if row["estimatedCost"] is not None else None
        )
        row["laborSavings"] = (
            row["savedHours"] * row["hourlyCost"]
            if row["savedHours"] is not None and row["hourlyCost"] is not None else None
        )
        benefits = [v for v in [row["laborSavings"], row["otherAnnualBenefit"]] if v is not None]
        row["annualBenefit"] = sum(benefits) if benefits else None
        row["paybackYears"] = (
            row["netCost"] / row["annualBenefit"]
            if row["netCost"] is not None and row["annualBenefit"] not in (None, 0) else None
        )

    action_columns = [
        "id", "area", "action", "owner", "dueDate", "priority", "status", "progress", "timing", "deliverableKpi",
        "dependencies", "nextAction", "updatedAt", "note", "source"
    ]
    actions = read_table(wb["アクション管理"], 5, action_columns, stop_when_blank_id=True)
    for row in actions:
        row["timing"] = None

    kpi_columns = [
        "id", "level", "area", "metric", "definition", "target", "actual", "unit", "direction", "status", "frequency", "source"
    ]
    kpis = read_table(wb["KGI_KPI"], 5, kpi_columns, stop_when_blank_id=True)
    for row in kpis:
        row["actual"] = None
        row["status"] = "未入力"

    payload = {
        "meta": {
            "title": "2027年度 経営改革ダッシュボード",
            "asOfDate": "2026-09-18",
            "targetYear": int(master["対象年度"]),
            "sourceWorkbook": SOURCE.name,
            "sourceUpdatedAt": "2026-09-18",
        },
        "master": {
            "bookingThresholds": thresholds,
            "currentGolfTax": master["利用税_現行"],
            "plannedGolfTaxCap": master["利用税_申請上限"],
            "annualFeeIncrease": master["年会費_改定額"],
            "memberPlayFeeIncrease": master["会員プレー料金_改定額"],
            "plannedRobotMowers": master["自動芝刈機_追加予定"],
            "bunkerReductionTarget": master["バンカー削減縮小目標"],
        },
        "prices": prices,
        "revenue": revenue,
        "courseQuality": course_quality,
        "workforce": workforce,
        "dx": dx,
        "investments": investments,
        "actions": actions,
        "kpis": kpis,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
